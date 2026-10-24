"""
Generador de reportes de auditoria (RF-14).

A partir de la lista de hallazgos ya enriquecida por el motor de riesgo,
construye el informe con tres secciones (documento seccion 11.6.4):
    1. Resumen ejecutivo.
    2. Matriz de riesgo / mapa de calor.
    3. Detalle de hallazgos.

Formatos de salida (caso de uso "Generar y descargar informe"):
    - PDF   : informe legible para presentar y archivar (libreria fpdf2).
    - EXCEL : hoja de calculo con los hallazgos filtrable/ordenable (openpyxl).
    - JSON  : datos crudos para integracion.

Los formatos PDF y Excel usan librerias puras de Python (fpdf2 y openpyxl), sin
dependencias del sistema operativo, de modo que funcionan directamente en
Windows. Si alguna libreria no esta instalada, ese formato se omite y se
notifica (flujo alternativo A1), sin interrumpir el resto.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from core.models import Hallazgo, Severidad
from riesgo.motor_riesgo import ResumenRiesgo, nivel_riesgo

# Orden de severidad para ordenar el detalle (mas grave primero).
_ORDEN_SEV = {
    Severidad.CRITICA: 0, Severidad.ALTA: 1, Severidad.MEDIA: 2,
    Severidad.BAJA: 3, Severidad.INFORMATIVA: 4,
}

# Colores por nivel de riesgo (seccion 5.3), como tuplas RGB.
_RGB_NIVEL = {
    "Critico": (192, 57, 43),
    "Alto": (230, 126, 34),
    "Medio": (201, 166, 0),
    "Bajo": (39, 174, 96),
}

# Colores por severidad (para las celdas del informe), como tuplas RGB.
_RGB_SEVERIDAD = {
    "Critica": (192, 57, 43),
    "Alta": (230, 126, 34),
    "Media": (201, 166, 0),
    "Baja": (46, 125, 50),
    "Informativa": (90, 107, 123),
}

_AZUL = (31, 58, 95)
_AZUL_CLARO = (44, 82, 130)
_GRIS_SUAVE = (238, 242, 247)


def _rgb_hex(rgb: tuple[int, int, int]) -> str:
    """Convierte una tupla RGB a hex sin '#', para openpyxl."""
    return "%02X%02X%02X" % rgb


def _lat(texto) -> str:
    """
    Normaliza el texto para las fuentes base del PDF (latin-1): sustituye los
    caracteres tipograficos comunes que no existen en latin-1 y descarta con
    seguridad cualquier otro, evitando que la generacion del PDF falle.
    """
    if texto is None:
        return ""
    s = str(texto)
    reemplazos = {
        "–": "-", "—": "-", "‘": "'", "’": "'",
        "“": '"', "”": '"', "•": "-", "…": "...",
        "×": "x", "→": "->", " ": " ",
    }
    for a, b in reemplazos.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


class GeneradorReportes:
    """Construye y guarda el informe de auditoria en PDF, Excel y JSON."""

    def __init__(self, directorio_salida: str | Path = "reportes_generados", logger=None):
        self.dir = Path(directorio_salida)
        self.dir.mkdir(exist_ok=True)
        self.logger = logger

    def generar(self, hallazgos: list[Hallazgo], resumen: ResumenRiesgo, meta: dict) -> dict:
        """
        Genera los formatos disponibles y devuelve un diccionario con las rutas
        creadas: {'pdf': ... | None, 'excel': ... | None, 'json': ...}.
        """
        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        nombre = meta.get("nombre_objetivo", "auditoria").replace(" ", "_")
        base = self.dir / f"informe_{nombre}_{marca}"

        h_ord = sorted(hallazgos, key=lambda x: (_ORDEN_SEV.get(x.severidad, 9), -(x.cvss or 0)))

        rutas = {"pdf": None, "excel": None, "json": None}
        rutas["json"] = self._generar_json(hallazgos, resumen, meta, base)
        rutas["pdf"] = self._generar_pdf(h_ord, resumen, meta, base)
        rutas["excel"] = self._generar_excel(h_ord, resumen, meta, base)
        return rutas

    # ------------------------------------------------------------------ JSON
    def _generar_json(self, hallazgos, resumen, meta, base) -> Path:
        datos = {
            "metadata": {
                "objetivo": meta.get("url", ""),
                "nombre": meta.get("nombre_objetivo", ""),
                "fecha": datetime.now().isoformat(timespec="seconds"),
                "total_hallazgos": len(hallazgos),
            },
            "resumen_riesgo": resumen.to_dict(),
            "hallazgos": [h.to_dict() for h in hallazgos],
        }
        ruta = base.with_suffix(".json")
        ruta.write_text(json.dumps(datos, indent=2, ensure_ascii=False), encoding="utf-8")
        self._log("info", "Informe JSON generado: %s", ruta)
        return ruta

    # ------------------------------------------------------------------ PDF
    def _generar_pdf(self, hallazgos, resumen: ResumenRiesgo, meta, base) -> Path | None:
        try:
            from fpdf import FPDF
            from fpdf.enums import XPos, YPos
        except Exception as exc:
            self._log("warning", "A1: no se pudo generar el PDF (falta fpdf2: %s).", exc)
            return None

        ruta = base.with_suffix(".pdf")
        try:
            pdf = FPDF(orientation="P", unit="mm", format="A4")
            pdf.set_auto_page_break(auto=True, margin=16)
            pdf.add_page()
            ancho = pdf.w - pdf.l_margin - pdf.r_margin

            # ---- Portada / cabecera ----
            pdf.set_fill_color(*_AZUL)
            pdf.rect(0, 0, pdf.w, 34, style="F")
            pdf.set_xy(pdf.l_margin, 9)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font("Helvetica", "B", 18)
            pdf.cell(ancho, 9, _lat("Informe de Auditoria de Seguridad"),
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_font("Helvetica", "", 11)
            pdf.cell(ancho, 7, _lat("Cooperativa de Ahorro y Credito Nuevo Amanecer, R.L."),
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(10)

            pdf.set_text_color(60, 60, 60)
            pdf.set_font("Helvetica", "", 11)
            objetivo = f"{meta.get('nombre_objetivo', '')} ({meta.get('url', '')})"
            pdf.cell(ancho, 7, _lat(f"Objetivo: {objetivo}"),
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.cell(ancho, 7, _lat(f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"),
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(4)

            # ---- 1. Resumen ejecutivo ----
            self._pdf_titulo(pdf, "1. Resumen ejecutivo", ancho)
            criticos = resumen.por_nivel.get("Critico", 0)
            altos = resumen.por_nivel.get("Alto", 0)
            texto = (
                f"La auditoria identifico {resumen.total} hallazgo(s), con un riesgo "
                f"promedio de {resumen.riesgo_promedio:.1f} (escala 1-25). "
            )
            if criticos or altos:
                texto += f"Atencion prioritaria: {criticos} critico(s) y {altos} alto(s)."
            else:
                texto += "No se detectaron hallazgos de nivel critico o alto."
            pdf.set_font("Helvetica", "", 11)
            pdf.set_text_color(40, 40, 40)
            pdf.multi_cell(ancho, 6, _lat(texto))
            pdf.ln(3)

            # Tarjetas de conteo por nivel.
            self._pdf_tarjetas_nivel(pdf, resumen, ancho)
            pdf.ln(4)

            # ---- 2. Matriz de riesgo ----
            self._pdf_titulo(pdf, "2. Matriz de riesgo (mapa de calor)", ancho)
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(80, 80, 80)
            pdf.multi_cell(ancho, 5, _lat(
                "Numero de hallazgos por combinacion de probabilidad (P) e impacto (I). "
                "El valor entre parentesis es el riesgo (P x I)."))
            pdf.ln(2)
            self._pdf_mapa_calor(pdf, resumen)
            pdf.ln(4)

            # ---- 3. Detalle de hallazgos ----
            self._pdf_titulo(pdf, "3. Detalle de hallazgos", ancho)
            if not hallazgos:
                pdf.set_font("Helvetica", "", 11)
                pdf.set_text_color(40, 40, 40)
                pdf.multi_cell(ancho, 6, _lat("No se generaron hallazgos."))
            else:
                for i, h in enumerate(hallazgos, 1):
                    self._pdf_hallazgo(pdf, i, h, ancho)

            pdf.output(str(ruta))
            self._log("info", "Informe PDF generado: %s", ruta)
            return ruta
        except Exception as exc:
            self._log("warning", "A1: no se pudo generar el PDF (%s).", exc)
            return None

    def _pdf_titulo(self, pdf, texto, ancho):
        from fpdf.enums import XPos, YPos
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(*_AZUL)
        pdf.cell(ancho, 9, _lat(texto), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        y = pdf.get_y()
        pdf.set_draw_color(*_AZUL_CLARO)
        pdf.line(pdf.l_margin, y, pdf.l_margin + ancho, y)
        pdf.ln(3)

    def _pdf_tarjetas_nivel(self, pdf, resumen, ancho):
        from fpdf.enums import XPos, YPos
        niveles = [("Total", resumen.total, _AZUL_CLARO)]
        for niv in ("Critico", "Alto", "Medio", "Bajo"):
            niveles.append((niv, resumen.por_nivel.get(niv, 0), _RGB_NIVEL[niv]))
        n = len(niveles)
        gap = 3
        w = (ancho - gap * (n - 1)) / n
        x0 = pdf.l_margin
        y0 = pdf.get_y()
        for idx, (lbl, valor, rgb) in enumerate(niveles):
            x = x0 + idx * (w + gap)
            pdf.set_xy(x, y0)
            pdf.set_fill_color(*rgb)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font("Helvetica", "B", 16)
            pdf.cell(w, 11, str(valor), border=0, align="C", fill=True,
                     new_x=XPos.LEFT, new_y=YPos.NEXT)
            pdf.set_xy(x, y0 + 11)
            pdf.set_font("Helvetica", "", 8)
            pdf.cell(w, 6, _lat(lbl.upper()), border=0, align="C", fill=True)
        pdf.set_xy(x0, y0 + 17)

    def _pdf_mapa_calor(self, pdf, resumen):
        from fpdf.enums import XPos, YPos
        celda = 22
        alto = 12
        etq = 14
        x0 = pdf.l_margin
        # Fila de encabezado (I=1..5).
        pdf.set_xy(x0, pdf.get_y())
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*_AZUL)
        pdf.set_fill_color(*_GRIS_SUAVE)
        pdf.cell(etq, alto, "", border=1, align="C", fill=True)
        for i in range(1, 6):
            pdf.cell(celda, alto, _lat(f"I={i}"), border=1, align="C", fill=True)
        pdf.ln(alto)
        # Filas P=5..1.
        for p in range(5, 0, -1):
            pdf.set_x(x0)
            pdf.set_font("Helvetica", "B", 9)
            pdf.set_text_color(*_AZUL)
            pdf.set_fill_color(*_GRIS_SUAVE)
            pdf.cell(etq, alto, _lat(f"P={p}"), border=1, align="C", fill=True)
            for i in range(1, 6):
                valor = p * i
                rgb = _RGB_NIVEL[nivel_riesgo(valor)]
                n = resumen.mapa_calor.get(p, {}).get(i, 0)
                pdf.set_fill_color(*rgb)
                pdf.set_text_color(255, 255, 255)
                pdf.set_font("Helvetica", "B", 10)
                texto = f"{n}  ({valor})" if n else f"({valor})"
                pdf.cell(celda, alto, _lat(texto), border=1, align="C", fill=True)
            pdf.ln(alto)
        pdf.set_text_color(40, 40, 40)

    def _pdf_hallazgo(self, pdf, i, h, ancho):
        from fpdf.enums import XPos, YPos
        rgb = _RGB_SEVERIDAD.get(h.severidad.value, (120, 120, 120))
        # Barra de titulo con la severidad.
        pdf.set_fill_color(*rgb)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 11)
        titulo = f"{i}. [{h.severidad.value}] {h.titulo or h.categoria.value}"
        pdf.multi_cell(ancho, 8, _lat(titulo), border=0, fill=True,
                       new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        # Metadatos.
        pdf.set_text_color(90, 90, 90)
        pdf.set_font("Helvetica", "", 9)
        riesgo = (h.probabilidad or 0) * (h.impacto or 0)
        meta = (f"Categoria: {h.categoria.value}   |   CVSS: {h.cvss:.1f}   |   "
                f"Riesgo: {h.nivel_riesgo} (P{h.probabilidad} x I{h.impacto} = {riesgo})   |   "
                f"Modulo: {h.modulo_origen}")
        pdf.multi_cell(ancho, 5, _lat(meta), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if h.ubicacion:
            pdf.multi_cell(ancho, 5, _lat(f"Ubicacion: {h.ubicacion}"),
                           new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        # Evidencia y recomendacion.
        pdf.set_text_color(40, 40, 40)
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(ancho, 5, _lat("Evidencia:"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(ancho, 5, _lat(h.evidencia), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(ancho, 5, _lat("Recomendacion:"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(ancho, 5, _lat(h.recomendacion), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(4)

    # ---------------------------------------------------------------- EXCEL
    def _generar_excel(self, hallazgos, resumen: ResumenRiesgo, meta, base) -> Path | None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except Exception as exc:
            self._log("warning", "A1: no se pudo generar el Excel (falta openpyxl: %s).", exc)
            return None

        ruta = base.with_suffix(".xlsx")
        try:
            wb = Workbook()
            azul = _rgb_hex(_AZUL)
            blanco = Font(color="FFFFFF", bold=True)
            centro = Alignment(horizontal="center", vertical="center")
            borde = Border(*[Side(style="thin", color="D9D9D9")] * 4)

            # ---- Hoja 1: Resumen ----
            ws = wb.active
            ws.title = "Resumen"
            ws["A1"] = "Informe de Auditoria de Seguridad"
            ws["A1"].font = Font(size=15, bold=True, color=azul)
            ws["A2"] = "Cooperativa de Ahorro y Credito Nuevo Amanecer, R.L."
            ws["A2"].font = Font(size=10, italic=True, color="666666")
            filas_meta = [
                ("Objetivo", meta.get("nombre_objetivo", "")),
                ("URL", meta.get("url", "")),
                ("Fecha", datetime.now().strftime("%d/%m/%Y %H:%M")),
                ("Total de hallazgos", resumen.total),
                ("Riesgo promedio (1-25)", round(resumen.riesgo_promedio, 1)),
            ]
            r = 4
            for k, v in filas_meta:
                ws.cell(row=r, column=1, value=k).font = Font(bold=True)
                ws.cell(row=r, column=2, value=v)
                r += 1

            r += 1
            ws.cell(row=r, column=1, value="Hallazgos por nivel de riesgo").font = Font(bold=True, color=azul)
            r += 1
            for niv in ("Critico", "Alto", "Medio", "Bajo"):
                c1 = ws.cell(row=r, column=1, value=niv)
                c1.fill = PatternFill("solid", fgColor=_rgb_hex(_RGB_NIVEL[niv]))
                c1.font = blanco
                ws.cell(row=r, column=2, value=resumen.por_nivel.get(niv, 0))
                r += 1

            r += 1
            ws.cell(row=r, column=1, value="Hallazgos por severidad").font = Font(bold=True, color=azul)
            r += 1
            for sev in ("Critica", "Alta", "Media", "Baja", "Informativa"):
                c1 = ws.cell(row=r, column=1, value=sev)
                c1.fill = PatternFill("solid", fgColor=_rgb_hex(_RGB_SEVERIDAD[sev]))
                c1.font = blanco
                ws.cell(row=r, column=2, value=resumen.por_severidad.get(sev, 0))
                r += 1
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 48

            # ---- Hoja 2: Hallazgos ----
            wh = wb.create_sheet("Hallazgos")
            columnas = ["#", "Severidad", "Hallazgo", "Categoria OWASP", "Modulo",
                        "CVSS", "Nivel de riesgo", "P", "I", "Ubicacion",
                        "Evidencia", "Recomendacion"]
            for col, titulo in enumerate(columnas, 1):
                c = wh.cell(row=1, column=col, value=titulo)
                c.fill = PatternFill("solid", fgColor=azul)
                c.font = blanco
                c.alignment = centro
                c.border = borde
            for idx, h in enumerate(hallazgos, 1):
                fila = idx + 1
                valores = [
                    idx, h.severidad.value, h.titulo or h.categoria.value,
                    h.categoria.value, h.modulo_origen,
                    round(h.cvss, 1) if h.cvss is not None else "",
                    h.nivel_riesgo, h.probabilidad, h.impacto,
                    h.ubicacion or "", h.evidencia, h.recomendacion,
                ]
                for col, valor in enumerate(valores, 1):
                    c = wh.cell(row=fila, column=col, value=valor)
                    c.border = borde
                    c.alignment = Alignment(vertical="top", wrap_text=col in (3, 10, 11, 12))
                # Colorear la celda de severidad.
                cs = wh.cell(row=fila, column=2)
                cs.fill = PatternFill("solid", fgColor=_rgb_hex(_RGB_SEVERIDAD.get(h.severidad.value, (120, 120, 120))))
                cs.font = blanco
                cs.alignment = centro
            anchos = [4, 12, 42, 26, 16, 7, 14, 4, 4, 30, 55, 55]
            for col, w in enumerate(anchos, 1):
                wh.column_dimensions[get_column_letter(col)].width = w
            wh.freeze_panes = "A2"
            wh.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(hallazgos) + 1}"

            # ---- Hoja 3: Mapa de riesgo ----
            wm = wb.create_sheet("Mapa de riesgo")
            wm.cell(row=1, column=1, value="P \\ I").fill = PatternFill("solid", fgColor=azul)
            wm.cell(row=1, column=1).font = blanco
            wm.cell(row=1, column=1).alignment = centro
            for i in range(1, 6):
                c = wm.cell(row=1, column=i + 1, value=f"I={i}")
                c.fill = PatternFill("solid", fgColor=azul); c.font = blanco; c.alignment = centro
            fila = 2
            for p in range(5, 0, -1):
                c = wm.cell(row=fila, column=1, value=f"P={p}")
                c.fill = PatternFill("solid", fgColor=azul); c.font = blanco; c.alignment = centro
                for i in range(1, 6):
                    valor = p * i
                    n = resumen.mapa_calor.get(p, {}).get(i, 0)
                    cc = wm.cell(row=fila, column=i + 1, value=f"{n} ({valor})")
                    cc.fill = PatternFill("solid", fgColor=_rgb_hex(_RGB_NIVEL[nivel_riesgo(valor)]))
                    cc.font = Font(color="FFFFFF", bold=True)
                    cc.alignment = centro
                fila += 1
            wm.column_dimensions["A"].width = 10
            for i in range(2, 7):
                wm.column_dimensions[get_column_letter(i)].width = 12

            wb.save(str(ruta))
            self._log("info", "Informe Excel generado: %s", ruta)
            return ruta
        except Exception as exc:
            self._log("warning", "A1: no se pudo generar el Excel (%s).", exc)
            return None

    # ------------------------------------------------------------------ util
    def _log(self, nivel, msg, *args):
        if self.logger:
            getattr(self.logger, nivel)(msg, *args)
